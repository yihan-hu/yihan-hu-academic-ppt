from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[1]
P=json.loads((ROOT/'references/brands/ki-template.json').read_text())
C=P['tokens']
def color(key):return RGBColor.from_string(C.get(key,key))
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def shape(s,x,y,w,h,fill,stroke=None):
 q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 q.fill.solid();q.fill.fore_color.rgb=color(fill)
 q.line.fill.background() if not stroke else None
 if stroke:q.line.color.rgb=color(stroke)
 return q
def txt(s,t,x,y,w,h,size=18,bold=False,fg='body',font='Aptos',align=None):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 f=q.text_frame;f.clear();f.word_wrap=True;f.margin_left=Inches(.01);f.margin_top=Inches(.01);f.margin_right=Inches(.01)
 p=f.paragraphs[0];p.text=t;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fg)
 if align:p.alignment=align
 return q
def base(section,title,index,subtitle=None):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color('paper')
 shape(s,.68,.32,.06,.26,'accent')
 txt(s,section.upper(),.87,.28,7,.30,10,True,'accent')
 txt(s,index,11.75,.28,.82,.27,10,False,'muted',align=PP_ALIGN.RIGHT)
 txt(s,title,.68,.86,11.9,.64,28,True,'heading')
 if subtitle:txt(s,subtitle,.69,1.52,11.6,.35,12,False,'muted')
 shape(s,.69,1.99,1.0,.025,'accent')
 shape(s,1.77,2.00,10.85,.009,'rule')
 shape(s,.69,7.05,11.92,.009,'rule')
 txt(s,'Karolinska Institutet  |  Academic PPT',.69,7.12,6.5,.22,8,False,'muted')
 return s
def notes(s,t):
 try:s.notes_slide.notes_text_frame.text=t
 except Exception:pass

s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color('cover')
logo=ROOT/'assets/brands/ki-logo-white.png'
s.shapes.add_picture(str(logo),Inches(10.7),Inches(.51),width=Inches(1.78))
txt(s,'Your research title',.69,1.84,10,.9,36,True,'white')
txt(s,'A concise scientific subtitle',.69,2.81,10,.60,22,False,'white')
shape(s,.69,4.03,3.0,.025,'white')
txt(s,'Author(s)',.70,4.30,8,.42,18,False,'white')
txt(s,'Karolinska Institutet',.70,4.84,8,.34,13,False,'white')
notes(s,'Replace placeholders. Keep the full deep-plum cover and the official logo; no decorative network graphic.')

s=base('Background','Previous evidence and research objective','02 / 07')
shape(s,.75,2.30,5.7,2.40,'panel');shape(s,6.79,2.30,5.76,2.40,'white')
txt(s,'PREVIOUS EVIDENCE',1.0,2.60,5.1,.32,11,True,'accent')
txt(s,'Published finding',1.0,3.08,5.1,.55,22,True,'heading')
txt(s,'Summarize prior findings with a reference.',1.0,3.81,5.1,.60,15)
txt(s,'CURRENT STUDY',7.05,2.60,5.1,.32,11,True,'accent')
txt(s,'Study question',7.05,3.08,5.1,.55,22,True,'heading')
txt(s,'Explain the methodological advance.',7.05,3.81,5.1,.60,15)
for i,t in enumerate(['Primary question','Secondary question','Clinical implications']):
 x=.80+i*4.1;txt(s,f'0{i+1}',x,5.35,.7,.4,22,True,'accent');txt(s,t,x+.7,5.38,3.25,.44,16,True,'heading')
notes(s,'Use source-backed text, and avoid redundant caveats.')

s=base('Methods','Study population and measurements','03 / 07', 'Study region  /  enrollment period')
shape(s,.76,2.35,5.65,1.39,'panel');shape(s,6.78,2.35,5.72,1.39,'white')
for x,n,title,desc in [(1.02,'N1','Cases','Eligibility / definition'),(7.02,'N2','Controls','Comparator groups')]:
 txt(s,n,x,2.57,1.65,.62,38,True,'cover');txt(s,title,x+1.9,2.56,3.4,.44,20,True,'heading');txt(s,desc,x,3.22,5.1,.32,12,False,'muted')
for i,(a,b) in enumerate([('Coverage','Population and catchment'),('Measurement','Exposure and assay'),('Follow-up','Outcomes and follow-up')]):
 x=.89+i*4.09;txt(s,f'0{i+1}',x,4.18,.6,.32,12,True,'accent');txt(s,a,x,4.57,3.8,.39,19,True,'heading');txt(s,b,x,5.12,3.7,.50,12,False,'muted')
notes(s,'Verify denominators, location, coverage, study dates, assay, and outcome definitions from the manuscript.')

s=base('Methods','Selected baseline characteristics','04 / 07')
headers=['Characteristic','Cases','Controls'];widths=[5.5,3.0,3.0];xs=[.83,6.33,9.3]
shape(s,.75,2.30,11.76,.56,'panel')
for x,w,h in zip(xs,widths,headers):txt(s,h,x,2.43,w,.30,13,True,'heading')
for i,r in enumerate(['Age, mean (SD)','Female, n (%)','Clinical feature, n (%)','Follow-up, median (IQR)']):
 y=3.09+i*.73
 if i%2:shape(s,.75,y-.06,11.76,.57,'white')
 txt(s,r,.83,y,5.2,.34,14)
 txt(s,'—',6.33,y,2.6,.34,14,False,'muted')
 txt(s,'—',9.3,y,2.6,.34,14,False,'muted')
notes(s,'These are intentionally empty placeholders, never copied from the ALS study.')

s=base('Results','Main effect estimate','05 / 07','Replace illustrative numbers with verified data')
shape(s,.76,2.35,5.62,3.50,'panel')
txt(s,'PRIMARY ESTIMATE',1.00,2.69,5.0,.32,11,True,'accent')
txt(s,'1.00',1.00,3.13,5.0,1.02,46,True,'cover')
txt(s,'95% CI 0.80–1.20',1.00,4.35,5.1,.45,19,True,'heading')
txt(s,'Contrast / unit of exposure',1.02,5.17,5.05,.40,14)
txt(s,'INTERPRETATION',6.94,2.69,5.1,.32,11,True,'accent')
txt(s,'One sentence summarizing the main estimate.',6.94,3.20,5.2,1.46,23,True,'heading')
notes(s,'Illustrative values only. Preserve effect type, CI or CrI, correct unit and denominators.')

s=base('Results','Associations across groups','06 / 07','Editable forest plot with illustrative data')
shape(s,.76,2.30,11.75,.55,'panel')
txt(s,'SUBGROUP',.92,2.44,4.5,.30,11,True,'accent')
txt(s,'EFFECT (95% CI)',9.10,2.44,3.15,.30,11,True,'accent')
pairs=[('Subgroup A',.81,1.10,1.42),('Subgroup B',.70,.94,1.25),('Subgroup C',.94,1.19,1.53),('Subgroup D',.77,1.02,1.37)]
for i,(name,lo,est,hi) in enumerate(pairs):
 y=3.10+i*.65;txt(s,name,.92,y,4,.32,14)
 x0=5.36
 def px(v):return x0+(v-.5)*2.25
 shape(s,px(lo),y+.13,max(.03,px(hi)-px(lo)),.025,'muted')
 q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(px(est)-.075),Inches(y+.07),Inches(.15),Inches(.15));q.fill.solid();q.fill.fore_color.rgb=color('accent');q.line.fill.background()
 txt(s,f'{est:.2f} ({lo:.2f}–{hi:.2f})',9.10,y,3.1,.32,13)
txt(s,'Reference = 1.0; replace effect values and rescale for each study.',.92,6.40,11.0,.37,11,False,'muted')
notes(s,'Editable marks and labels. All illustrative values are placeholders.')

s=base('Conclusions','Take-home messages','07 / 07')
for i,t in enumerate(['Primary conclusion in one sentence.','Secondary conclusion in one sentence.','Implication or next scientific question.']):
 y=2.60+i*1.18;txt(s,f'0{i+1}',1.10,y,.92,.48,25,True,'accent');txt(s,t,2.35,y,9.72,.53,20,True,'heading')
 if i<2:shape(s,2.35,y+.86,9.8,.008,'rule')
txt(s,'Thank you',.89,6.48,6,.45,19,True,'heading')
notes(s,'Three concise takeaways; no study-specific numerical claims.')

out=Path(__file__).resolve().parents[1]/'assets/templates/KI-template-starter.pptx'
out.parent.mkdir(parents=True,exist_ok=True)
prs.save(str(out));print(out)
